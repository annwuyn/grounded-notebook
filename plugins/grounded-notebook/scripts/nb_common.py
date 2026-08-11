"""nb_common — chuyển tài liệu sang Markdown, chia chunk giữ offset ký tự.

Mọi chunk giữ (start, end) là offset trong chuỗi Markdown đã chuyển đổi.
Nhờ đó reader có thể tô sáng đúng đoạn trong toàn văn mà không cần đoán.
"""
import io
import json
import os
import re
import sys
import unicodedata
from bisect import bisect_right

TARGET_WORDS = 80
MIN_WORDS = 55
MAX_WORDS = 120

TEXT_EXT = {".md", ".markdown", ".txt", ".text", ".rst", ".org"}
SUPPORTED = TEXT_EXT | {
    ".pdf", ".docx", ".pptx", ".xlsx", ".xlsm", ".csv", ".tsv",
    ".html", ".htm", ".json",
}


# ---------------------------------------------------------------- utilities

def _soft(name):
    try:
        return __import__(name)
    except Exception:
        return None


def warn(msg):
    print("  ! " + msg, file=sys.stderr)


def norm_ws(s):
    return re.sub(r"\s+", " ", s).strip()


def fold(s):
    """Chuẩn hoá để so khớp trích dẫn: bỏ dấu nháy cong, gộp khoảng trắng."""
    s = unicodedata.normalize("NFC", s)
    s = (s.replace("\u2018", "'").replace("\u2019", "'")
          .replace("\u201c", '"').replace("\u201d", '"')
          .replace("\u2013", "-").replace("\u2014", "-")
          .replace("\u00a0", " ").replace("\u200b", ""))
    return re.sub(r"\s+", " ", s).strip()


def slugify(s, maxlen=40):
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = re.sub(r"[^A-Za-z0-9]+", "-", s).strip("-").lower()
    return (s or "doc")[:maxlen]


# ------------------------------------------------------------- md builder

class MD:
    """Gom Markdown và ghi lại khoảng offset của từng 'trang'."""

    def __init__(self):
        self.buf = []
        self.n = 0
        self.pages = []          # [(page_no, start, end)]
        self._page_open = None

    def add(self, text):
        if not text:
            return
        self.buf.append(text)
        self.n += len(text)

    def para(self, text):
        text = text.rstrip()
        if not text:
            return
        self.add(text + "\n\n")

    def open_page(self, no):
        self.close_page()
        self._page_open = (no, self.n)

    def close_page(self):
        if self._page_open is not None:
            no, start = self._page_open
            if self.n > start:
                self.pages.append((no, start, self.n))
            self._page_open = None

    def text(self):
        self.close_page()
        return "".join(self.buf)


# --------------------------------------------------------------- converters

def _conv_pdf(path, md):
    pdfplumber = _soft("pdfplumber")
    if pdfplumber:
        with pdfplumber.open(path) as pdf:
            for i, page in enumerate(pdf.pages, 1):
                md.open_page(i)
                txt = page.extract_text() or ""
                md.para(_clean_pdf_text(txt))
        md.close_page()
        if md.n > 0:
            return
        warn("pdfplumber không lấy được text, thử pypdf")
    pypdf = _soft("pypdf")
    if not pypdf:
        raise RuntimeError("Cần pdfplumber hoặc pypdf: pip install pdfplumber")
    reader = pypdf.PdfReader(path)
    for i, page in enumerate(reader.pages, 1):
        md.open_page(i)
        md.para(_clean_pdf_text(page.extract_text() or ""))
    md.close_page()


def _clean_pdf_text(txt):
    """Nối lại dòng bị ngắt giữa câu, giữ nguyên ngắt đoạn."""
    txt = txt.replace("\u00ad", "")
    lines = [l.rstrip() for l in txt.split("\n")]
    out, cur = [], []
    for l in lines:
        if not l.strip():
            if cur:
                out.append(" ".join(cur))
                cur = []
            out.append("")
            continue
        if cur and re.search(r"[a-zà-ỹ0-9,;:\-–]$", cur[-1], re.I) \
           and not re.match(r"^\s*[-•*\d]", l):
            if cur[-1].endswith("-"):
                cur[-1] = cur[-1][:-1] + l.lstrip()
            else:
                cur.append(l.strip())
        else:
            if cur:
                out.append(" ".join(cur))
            cur = [l.strip()]
    if cur:
        out.append(" ".join(cur))
    body = "\n\n".join(p for p in "\n".join(out).split("\n\n"))
    return re.sub(r"\n{3,}", "\n\n", body).strip()


def _conv_docx(path, md):
    docx = _soft("docx")
    if not docx:
        raise RuntimeError("Cần python-docx: pip install python-docx")
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    doc = docx.Document(path)
    body = doc.element.body

    def iter_blocks():
        for child in body.iterchildren():
            tag = child.tag.split("}")[-1]
            if tag == "p":
                yield Paragraph(child, doc)
            elif tag == "tbl":
                yield Table(child, doc)

    for block in iter_blocks():
        if isinstance(block, Paragraph):
            t = block.text.strip()
            if not t:
                continue
            style = (block.style.name or "").lower()
            m = re.match(r"heading (\d)", style)
            if m:
                md.para("#" * min(6, int(m.group(1))) + " " + t)
            elif "title" in style:
                md.para("# " + t)
            elif "list" in style:
                md.para("- " + t)
            else:
                md.para(t)
        else:
            rows = []
            for r in block.rows:
                rows.append([norm_ws(c.text) for c in r.cells])
            if not rows:
                continue
            head = rows[0]
            md.add("| " + " | ".join(head) + " |\n")
            md.add("|" + "---|" * len(head) + "\n")
            for r in rows[1:]:
                r = (r + [""] * len(head))[:len(head)]
                md.add("| " + " | ".join(r) + " |\n")
            md.add("\n")


def _conv_pptx(path, md):
    pptx = _soft("pptx")
    if not pptx:
        raise RuntimeError("Cần python-pptx: pip install python-pptx")
    prs = pptx.Presentation(path)
    for i, slide in enumerate(prs.slides, 1):
        md.open_page(i)
        md.para("## Slide %d" % i)
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    t = "".join(r.text for r in p.runs).strip()
                    if t:
                        md.para(t)
            if getattr(shape, "has_table", False):
                for r in shape.table.rows:
                    md.para("| " + " | ".join(
                        norm_ws(c.text) for c in r.cells) + " |")
        if slide.has_notes_slide:
            nt = (slide.notes_slide.notes_text_frame.text or "").strip()
            if nt:
                md.para("> Ghi chú: " + nt)
    md.close_page()


def _conv_xlsx(path, md):
    op = _soft("openpyxl")
    if not op:
        raise RuntimeError("Cần openpyxl: pip install openpyxl")
    wb = op.load_workbook(path, data_only=True, read_only=True)
    for si, ws in enumerate(wb.worksheets, 1):
        md.open_page(si)
        md.para("## " + ws.title)
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        width = max(len(r) for r in rows)
        for ri, r in enumerate(rows):
            cells = ["" if v is None else norm_ws(str(v))
                     for v in (list(r) + [None] * (width - len(r)))]
            md.add("| " + " | ".join(cells) + " |\n")
            if ri == 0:
                md.add("|" + "---|" * width + "\n")
        md.add("\n")
    md.close_page()


def _conv_html(path, md):
    bs4 = _soft("bs4")
    raw = open(path, "r", encoding="utf-8", errors="replace").read()
    if not bs4:
        md.para(re.sub(r"<[^>]+>", " ", raw))
        return
    soup = bs4.BeautifulSoup(raw, "html.parser")
    for bad in soup(["script", "style", "nav", "footer", "noscript"]):
        bad.decompose()
    for el in soup.find_all(["h1", "h2", "h3", "h4", "h5", "h6",
                            "p", "li", "blockquote", "pre", "td"]):
        t = norm_ws(el.get_text(" "))
        if not t:
            continue
        if el.name.startswith("h") and len(el.name) == 2:
            md.para("#" * int(el.name[1]) + " " + t)
        elif el.name == "li":
            md.para("- " + t)
        elif el.name == "blockquote":
            md.para("> " + t)
        else:
            md.para(t)


def _conv_csv(path, md):
    import csv
    delim = "\t" if path.lower().endswith(".tsv") else ","
    with open(path, newline="", encoding="utf-8", errors="replace") as f:
        rows = list(csv.reader(f, delimiter=delim))
    if not rows:
        return
    width = max(len(r) for r in rows)
    for ri, r in enumerate(rows):
        cells = [norm_ws(c) for c in (r + [""] * (width - len(r)))]
        md.add("| " + " | ".join(cells) + " |\n")
        if ri == 0:
            md.add("|" + "---|" * width + "\n")
    md.add("\n")


def _conv_text(path, md):
    md.add(open(path, "r", encoding="utf-8", errors="replace").read())


def _conv_markitdown(path, md):
    mid = _soft("markitdown")
    if not mid:
        raise RuntimeError("Định dạng không hỗ trợ và chưa có markitdown")
    res = mid.MarkItDown().convert(path)
    md.add(res.text_content or "")


def to_markdown(path):
    """→ (md_text, pages, kind). pages = [(no, start, end)] nếu có phân trang."""
    ext = os.path.splitext(path)[1].lower()
    md = MD()
    if ext == ".pdf":
        _conv_pdf(path, md); kind = "pdf"
    elif ext == ".docx":
        _conv_docx(path, md); kind = "docx"
    elif ext == ".pptx":
        _conv_pptx(path, md); kind = "pptx"
    elif ext in (".xlsx", ".xlsm"):
        _conv_xlsx(path, md); kind = "xlsx"
    elif ext in (".html", ".htm"):
        _conv_html(path, md); kind = "html"
    elif ext in (".csv", ".tsv"):
        _conv_csv(path, md); kind = "csv"
    elif ext in TEXT_EXT:
        _conv_text(path, md); kind = "text"
    elif ext == ".json":
        md.add("```json\n" + open(path, encoding="utf-8",
                                  errors="replace").read() + "\n```")
        kind = "json"
    else:
        _conv_markitdown(path, md); kind = "other"
    text = md.text()
    text = re.sub(r"[ \t]+\n", "\n", text)
    return text, md.pages, kind


# ---------------------------------------------------------------- chunking

_BLOCK_RE = re.compile(r"[^\n]+(?:\n[^\n]+)*")
_SENT_RE = re.compile(r"[^.!?…]*[.!?…]+[\"')\]]*\s*|[^.!?…]+$")
_HEAD_RE = re.compile(r"^(#{1,6})\s+(.*)$")


def _pieces(md_text):
    """Sinh (start, end, kind, heading_path) cho từng khối văn bản."""
    heads = []          # [(level, text)]
    for m in _BLOCK_RE.finditer(md_text):
        block = m.group(0)
        hm = _HEAD_RE.match(block.strip())
        if hm and "\n" not in block.strip():
            lvl = len(hm.group(1))
            heads = [h for h in heads if h[0] < lvl]
            heads.append((lvl, hm.group(2).strip()))
            yield m.start(), m.end(), "heading", [t for _, t in heads]
            continue
        path = [t for _, t in heads]
        words = len(block.split())
        if words <= MAX_WORDS:
            yield m.start(), m.end(), "block", path
            continue
        base = m.start()
        for sm in _SENT_RE.finditer(block):
            if sm.group(0).strip():
                yield base + sm.start(), base + sm.end(), "block", path


def _page_of(pages, starts, off):
    if not pages:
        return None
    i = bisect_right(starts, off) - 1
    if i < 0:
        i = 0
    return pages[i][0]


def chunk_document(md_text, pages, sid):
    """→ list chunk dict: cid, sid, page, heading, start, end, order."""
    starts = [p[1] for p in pages]
    chunks = []
    cur = []          # [(start, end)]
    cur_head = []
    cur_words = 0

    def flush():
        nonlocal cur, cur_words
        if not cur:
            return
        s, e = cur[0][0], cur[-1][1]
        page = _page_of(pages, starts, s)
        chunks.append({
            "sid": sid, "page": page, "heading": " › ".join(cur_head[-2:]),
            "start": s, "end": e,
        })
        cur, cur_words = [], 0

    for s, e, kind, head in _pieces(md_text):
        if kind == "heading":
            flush()
            cur_head = head
            continue
        p_page = _page_of(pages, starts, s)
        if cur and pages and _page_of(pages, starts, cur[0][0]) != p_page:
            flush()
        cur_head = head or cur_head
        cur.append((s, e))
        cur_words += len(md_text[s:e].split())
        if cur_words >= TARGET_WORDS:
            flush()
    flush()

    # gộp chunk quá ngắn vào chunk trước nếu cùng trang
    merged = []
    for c in chunks:
        if merged and len(md_text[c["start"]:c["end"]].split()) < 12 \
           and merged[-1]["page"] == c["page"] \
           and merged[-1]["heading"] == c["heading"]:
            merged[-1]["end"] = c["end"]
        else:
            merged.append(c)

    per_page = {}
    for i, c in enumerate(merged):
        c["order"] = i
        if c["page"] is not None:
            per_page[c["page"]] = per_page.get(c["page"], 0) + 1
            c["cid"] = "%s:p%d.%d" % (sid, c["page"], per_page[c["page"]])
        else:
            c["cid"] = "%s:c%d" % (sid, i + 1)
    return merged


# ------------------------------------------------------------------- index

def load_index(path):
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def save_index(path, idx):
    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(idx, f, ensure_ascii=False, indent=1)


def read_doc(idx_path, src):
    """Đọc markdown đã chuyển đổi của một nguồn."""
    base = os.path.dirname(os.path.abspath(idx_path))
    p = os.path.join(base, src["markdown"])
    with open(p, encoding="utf-8") as f:
        return f.read()


CID_RE = re.compile(r"S\d+:(?:p\d+\.\d+|c\d+)")
MARKER_RE = re.compile(
    r"\[(S\d+:(?:p\d+\.\d+|c\d+)(?:\s*;\s*S\d+:(?:p\d+\.\d+|c\d+))*)\]")
LOOSE_MARKER_RE = re.compile(r"\[\s*S\s*\d+\s*[:.]?[^\]\n]{0,40}\]")
